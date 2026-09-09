from pathlib import Path
import tempfile
from io import BytesIO

from dotenv import load_dotenv

from fastapi import (
    FastAPI,
    UploadFile,
    File
)

from fastapi.responses import (
    HTMLResponse,
    PlainTextResponse,
    StreamingResponse
)

from fastapi.staticfiles import StaticFiles

from pydantic import BaseModel, Field

from docx import Document

from .summarizer import TextSummarizer
from .file_extractor import extract_text


# ==========================================
# LOAD ENVIRONMENT
# ==========================================

load_dotenv()


# ==========================================
# FASTAPI APP
# ==========================================

app = FastAPI(
    title="AI Text Summarizer",
    description=(
        "AI Text Summarizer using "
        "FastAPI, LangChain and Groq"
    ),
    version="3.0.0"
)


# ==========================================
# BASE DIRECTORY
# ==========================================

BASE_DIR = Path(__file__).resolve().parent


# ==========================================
# STATIC FILES
# ==========================================

app.mount(
    "/static",
    StaticFiles(
        directory=BASE_DIR / "static"
    ),
    name="static"
)


# ==========================================
# SUMMARIZER
# ==========================================

summarizer = TextSummarizer()


# ==========================================
# REQUEST MODELS
# ==========================================

class SummarizeRequest(BaseModel):

    text: str = Field(
        ...,
        min_length=20,
        max_length=100000
    )

    style: str = "simple"

    length: str = "medium"


class DownloadRequest(BaseModel):

    summary: str


# ==========================================
# ALLOWED VALUES
# ==========================================

ALLOWED_STYLES = {
    "simple",
    "professional",
    "bullet_points"
}


ALLOWED_LENGTHS = {
    "short",
    "medium",
    "detailed"
}


ALLOWED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx"
}


# ==========================================
# HOME PAGE
# ==========================================

@app.get(
    "/",
    response_class=HTMLResponse
)
async def home():

    html_file = (
        BASE_DIR
        / "templates"
        / "index.html"
    )

    return HTMLResponse(
        html_file.read_text(
            encoding="utf-8"
        )
    )


# ==========================================
# SUMMARIZE NORMAL TEXT
# ==========================================

@app.post("/api/summarize")
async def summarize(
    request: SummarizeRequest
):

    # --------------------------------------
    # Validate style
    # --------------------------------------

    if request.style not in ALLOWED_STYLES:

        return {
            "success": False,
            "error": "Invalid summary style."
        }


    # --------------------------------------
    # Validate length
    # --------------------------------------

    if request.length not in ALLOWED_LENGTHS:

        return {
            "success": False,
            "error": "Invalid summary length."
        }


    try:

        # ----------------------------------
        # Generate summary
        # ----------------------------------

        result = summarizer.summarize(
            text=request.text,
            style=request.style,
            length=request.length
        )


        # ----------------------------------
        # Return response
        # ----------------------------------

        return {
            "success": True,
            "summary": result["summary"],
            "chunks": result["chunks"],
            "characters": len(request.text),
            "words": len(request.text.split())
        }


    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }


# ==========================================
# SUMMARIZE UPLOADED FILE
# ==========================================

@app.post("/api/summarize-file")
async def summarize_file(
    file: UploadFile = File(...),
    style: str = "simple",
    length: str = "medium"
):

    # --------------------------------------
    # Get filename safely
    # --------------------------------------

    filename = file.filename or ""


    # --------------------------------------
    # Get extension
    # --------------------------------------

    extension = Path(
        filename
    ).suffix.lower()


    # --------------------------------------
    # Validate extension
    # --------------------------------------

    if extension not in ALLOWED_EXTENSIONS:

        return {
            "success": False,
            "error": (
                "Only TXT, PDF and DOCX "
                "files are supported."
            )
        }


    # --------------------------------------
    # Validate style
    # --------------------------------------

    if style not in ALLOWED_STYLES:

        return {
            "success": False,
            "error": "Invalid summary style."
        }


    # --------------------------------------
    # Validate length
    # --------------------------------------

    if length not in ALLOWED_LENGTHS:

        return {
            "success": False,
            "error": "Invalid summary length."
        }


    temp_path = None


    try:

        # ----------------------------------
        # Read uploaded file
        # ----------------------------------

        file_content = await file.read()


        if not file_content:

            return {
                "success": False,
                "error": "Uploaded file is empty."
            }


        # ----------------------------------
        # Create temporary file
        # ----------------------------------

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=extension
        ) as temp_file:

            temp_file.write(
                file_content
            )

            temp_path = temp_file.name


        # ----------------------------------
        # Extract text
        # ----------------------------------

        extracted_text = extract_text(
            temp_path
        )


        # ----------------------------------
        # Check extracted text
        # ----------------------------------

        if not extracted_text.strip():

            return {
                "success": False,
                "error": (
                    "No readable text found "
                    "in the uploaded file."
                )
            }


        # ----------------------------------
        # Generate summary
        # ----------------------------------

        result = summarizer.summarize(
            text=extracted_text,
            style=style,
            length=length
        )


        # ----------------------------------
        # Return response
        # ----------------------------------

        return {
            "success": True,
            "filename": filename,
            "characters": len(extracted_text),
            "words": len(extracted_text.split()),
            "chunks": result["chunks"],
            "summary": result["summary"]
        }


    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }


    finally:

        # ----------------------------------
        # Delete temporary file
        # ----------------------------------

        if temp_path:

            try:

                Path(
                    temp_path
                ).unlink(
                    missing_ok=True
                )

            except Exception:

                pass


# ==========================================
# DOWNLOAD TXT
# ==========================================

@app.post("/api/download/txt")
async def download_txt(
    data: DownloadRequest
):

    summary = data.summary.strip()


    # --------------------------------------
    # Validate summary
    # --------------------------------------

    if not summary:

        return {
            "success": False,
            "error": "Summary is empty."
        }


    # --------------------------------------
    # Return TXT file
    # --------------------------------------

    return PlainTextResponse(
        content=summary,
        headers={
            "Content-Disposition":
                'attachment; filename="summary.txt"'
        }
    )


# ==========================================
# DOWNLOAD DOCX
# ==========================================

@app.post("/api/download/docx")
async def download_docx(
    data: DownloadRequest
):

    summary = data.summary.strip()


    # --------------------------------------
    # Validate summary
    # --------------------------------------

    if not summary:

        return {
            "success": False,
            "error": "Summary is empty."
        }


    # --------------------------------------
    # Create DOCX document
    # --------------------------------------

    document = Document()


    # --------------------------------------
    # Add heading
    # --------------------------------------

    document.add_heading(
        "AI Generated Summary",
        level=1
    )


    # --------------------------------------
    # Add summary paragraphs
    # --------------------------------------

    for paragraph in summary.split("\n"):

        paragraph = paragraph.strip()


        if paragraph:

            document.add_paragraph(
                paragraph
            )


    # --------------------------------------
    # Save document to memory
    # --------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)


    # --------------------------------------
    # Return DOCX file
    # --------------------------------------

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="summary.docx"'
        }
    )